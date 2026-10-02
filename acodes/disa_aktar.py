## Tabloları Excel (.xlsx) veya CSV olarak dışa aktarma ##
# Tablo ekranda nasıl görünüyorsa (arama, filtre, sıralama) öyle aktarılır.

import csv
import datetime
import os
import re

from PyQt5.QtCore import Qt
from PyQt5.QtWidgets import QFileDialog, QMenu, QMessageBox

from database.yedek import yedek_klasoru

EXCEL = "Excel (*.xlsx)"
CSV = "CSV - Excel'de açılır (*.csv)"
_TARIH = re.compile(r"(\d{2})\.(\d{2})\.(\d{4})")


def tablo_verisi(tablo):
    """(başlıklar, satırlar): ekrandaki sırayla, gizli kolonlar ve tamamen boş satırlar hariç."""
    kolonlar = [c for c in range(tablo.columnCount()) if not tablo.isColumnHidden(c)]
    basliklar = []
    for c in kolonlar:
        baslik = tablo.horizontalHeaderItem(c)
        basliklar.append(baslik.text() if baslik else "")
    satirlar = []
    for r in range(tablo.rowCount()):
        if tablo.isRowHidden(r):
            continue
        satir = [tablo.item(r, c).text() if tablo.item(r, c) else "" for c in kolonlar]
        if any(satir):
            satirlar.append(satir)
    return basliklar, satirlar


def excel_degeri(metin):
    """Kısa tam sayılar sayı, gg.aa.yyyy tarih olarak yazılır. Telefon gibi uzun
    veya 0 ile başlayan rakam dizileri metin kalır (baştaki sıfır kaybolmasın)."""
    if metin.isdigit() and len(metin) <= 6 and not (len(metin) > 1 and metin.startswith("0")):
        return int(metin)
    tarih = _TARIH.fullmatch(metin)
    if tarih:
        gun, ay, yil = map(int, tarih.groups())
        try:
            return datetime.date(yil, ay, gun)
        except ValueError:
            pass
    return metin


def csv_yaz(yol, basliklar, satirlar):
    # utf-8-sig ve ";" ile Türkçe Excel dosyayı karakterleri bozmadan, kolonlara ayırarak açar
    with open(yol, "w", newline="", encoding="utf-8-sig") as dosya:
        yazici = csv.writer(dosya, delimiter=";")
        yazici.writerow(basliklar)
        yazici.writerows(satirlar)


def excel_yaz(yol, basliklar, satirlar, sayfa_adi="Liste"):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    kitap = Workbook()
    sayfa = kitap.active
    # Excel sayfa adında \ / * ? : [ ] kullanılamaz ve en fazla 31 karakter olabilir
    sayfa.title = " ".join(re.sub(r"[\\/*?:\[\]]", " ", sayfa_adi).split())[:31] or "Liste"
    sayfa.append(basliklar)
    for hucre in sayfa[1]:
        hucre.font = Font(bold=True)
        hucre.fill = PatternFill("solid", fgColor="FFFF99")
    for satir in satirlar:
        sayfa.append([excel_degeri(d) for d in satir])
    for c, baslik in enumerate(basliklar, start=1):
        en_uzun = max([len(baslik)] + [len(s[c - 1]) for s in satirlar])
        sayfa.column_dimensions[get_column_letter(c)].width = min(max(en_uzun + 2, 8), 60)
        for hucre in sayfa.iter_rows(min_row=2, min_col=c, max_col=c):
            if isinstance(hucre[0].value, datetime.date):
                hucre[0].number_format = "DD.MM.YYYY"
    sayfa.freeze_panes = "A2"
    if satirlar:
        sayfa.auto_filter.ref = sayfa.dimensions
    kitap.save(yol)


def disa_aktar(parent, tablo, ad):
    """Dosya adı sorar ve tabloyu seçilen biçimde kaydeder. Kaydedilen yolu (veya None) döndürür."""
    basliklar, satirlar = tablo_verisi(tablo)
    if not satirlar:
        QMessageBox.information(parent, "Bilgi", "Aktarılacak kayıt yok. Önce listeyi görüntüleyin.")
        return None
    temiz_ad = re.sub(r"\W+", "_", ad).strip("_")
    dosya_adi = f"{temiz_ad}_{datetime.date.today():%Y%m%d}.xlsx"
    baslangic = os.path.join(os.path.expanduser("~/Documents") if os.path.isdir(os.path.expanduser("~/Documents"))
                             else yedek_klasoru(), dosya_adi)
    yol, secilen = QFileDialog.getSaveFileName(parent, f"Dışa Aktar: {ad}", baslangic, f"{EXCEL};;{CSV}")
    if not yol:
        return None
    csv_mi = secilen == CSV or yol.lower().endswith(".csv")
    uzanti = ".csv" if csv_mi else ".xlsx"
    if not yol.lower().endswith(uzanti):
        yol = os.path.splitext(yol)[0] + uzanti
    try:
        if csv_mi:
            csv_yaz(yol, basliklar, satirlar)
        else:
            excel_yaz(yol, basliklar, satirlar, ad)
    except ImportError:
        QMessageBox.warning(parent, "Uyarı!", "Excel dosyası için openpyxl kurulu değil.\n"
                                              "CSV olarak kaydedebilir veya 'pip install openpyxl' ile kurabilirsiniz.")
        return None
    except OSError as hata:
        QMessageBox.warning(parent, "Uyarı!", f"Dosya kaydedilemedi:\n{hata}")
        return None
    QMessageBox.information(parent, "Bilgi", f"{len(satirlar)} satır kaydedildi:\n{yol}")
    return yol


def sag_tik_menusu(parent, tablo, ad):
    """Tabloya sağ tıklanınca 'Dışa aktar...' seçeneği çıkar. Paneller tablo.sag_tik_eylemleri listesine
    satıra özel işlemler ekleyebilir: her biri satır sırası alıp [(metin, işlev, açık mı)] döndüren fonksiyon
    (ör. kitap satırında Düzenle / Ödünç ver / İade al). Sağ tıklanan satır seçilir."""
    tablo.setContextMenuPolicy(Qt.CustomContextMenu)
    tablo.sag_tik_eylemleri = []
    tablo.sag_tik_menu = lambda konum: menu_kur(parent, tablo, ad, konum)
    tablo.customContextMenuRequested.connect(lambda konum: tablo.sag_tik_menu(konum).exec_(
        tablo.viewport().mapToGlobal(konum)))


def menu_kur(parent, tablo, ad, konum):
    menu = QMenu(tablo)
    indeks = tablo.indexAt(konum)
    if indeks.isValid() and tablo.sag_tik_eylemleri:
        tablo.selectRow(indeks.row())
        for kaynak in tablo.sag_tik_eylemleri:
            for metin, islev, acik in kaynak(indeks.row()):
                menu.addAction(metin, islev).setEnabled(acik)
        if not menu.isEmpty():
            menu.addSeparator()
    menu.addAction("Excel / CSV olarak dışa aktar...", lambda: disa_aktar(parent, tablo, ad))
    return menu
