## Excel / CSV dışa aktarma testleri ##
import csv
import datetime

import pytest
from openpyxl import load_workbook
from PySide6.QtCore import Qt
from PySide6.QtWidgets import QFileDialog, QTableWidget

from acodes import disa_aktar as da
from acodes.guest import Guest
from acodes.library import Library
from acodes.tablo import tablo_ayarla, tabloya_yaz


@pytest.mark.parametrize("metin,beklenen", [
    ("123", 123), ("0", 0),
    ("5551112233", "5551112233"),       # telefon metin kalır
    ("0123", "0123"),                   # baştaki sıfır korunur
    ("05.01.2026", datetime.date(2026, 1, 5)),
    ("31.02.2026", "31.02.2026"),       # geçersiz tarih metin kalır
    ("Anne'nin Günlüğü", "Anne'nin Günlüğü"), ("", ""),
])
def test_excel_degeri(metin, beklenen):
    assert da.excel_degeri(metin) == beklenen


@pytest.fixture
def tablo(app):
    t = QTableWidget(0, 3)
    t.setHorizontalHeaderLabels(["Adı", "Yılı", "Tarih"])
    tablo_ayarla(t)
    tabloya_yaz(t, [["Şiir", "2005", "05.01.2026"], ["Çanlar", "1992", "01.12.2025"], ["", "", ""]])
    return t


def test_tablo_verisi_ekrandaki_sirayla(tablo):
    tablo.sortItems(1, Qt.AscendingOrder)
    basliklar, satirlar = da.tablo_verisi(tablo)
    assert basliklar == ["Adı", "Yılı", "Tarih"]
    assert satirlar == [["Çanlar", "1992", "01.12.2025"], ["Şiir", "2005", "05.01.2026"]]  # boş satır yok


def test_csv_turkce_excel_icin(tmp_path, tablo):
    yol = tmp_path / "liste.csv"
    da.csv_yaz(str(yol), *da.tablo_verisi(tablo))
    assert yol.read_bytes().startswith(b"\xef\xbb\xbf")          # UTF-8 BOM
    with open(yol, encoding="utf-8-sig", newline="") as f:
        assert list(csv.reader(f, delimiter=";"))[1] == ["Şiir", "2005", "05.01.2026"]


def test_excel_bicimleri(tmp_path, tablo):
    yol = tmp_path / "liste.xlsx"
    da.excel_yaz(str(yol), *da.tablo_verisi(tablo), sayfa_adi="İstatistik: Tür/Yıl")
    sayfa = load_workbook(yol).active
    assert sayfa.title == "İstatistik Tür Yıl"                    # Excel'de yasak karakterler atılır
    assert [h.value for h in sayfa[1]] == ["Adı", "Yılı", "Tarih"] and sayfa["A1"].font.bold
    assert sayfa["B2"].value == 2005
    assert sayfa["C2"].value.date() == datetime.date(2026, 1, 5) and sayfa["C2"].number_format == "DD.MM.YYYY"
    assert sayfa.freeze_panes == "A2" and sayfa.auto_filter.ref == "A1:C3"


def dosya_sec(monkeypatch, yol, filtre=da.EXCEL):
    monkeypatch.setattr(QFileDialog, "getSaveFileName", staticmethod(lambda *a: (str(yol), filtre)))


def test_disa_aktar_uzanti_ekler(app, uyarilar, monkeypatch, tmp_path, tablo):
    dosya_sec(monkeypatch, tmp_path / "rapor", da.CSV)
    assert da.disa_aktar(None, tablo, "Deneme") == str(tmp_path / "rapor.csv")
    dosya_sec(monkeypatch, tmp_path / "rapor.txt", da.EXCEL)
    assert da.disa_aktar(None, tablo, "Deneme") == str(tmp_path / "rapor.xlsx")
    assert uyarilar[-1].startswith("2 satır kaydedildi")


def test_bos_tablo_aktarilmaz(app, uyarilar, monkeypatch):
    monkeypatch.setattr(QFileDialog, "getSaveFileName", staticmethod(lambda *a: pytest.fail("dosya sorulmamalı")))
    assert da.disa_aktar(None, QTableWidget(1, 2), "Boş") is None
    assert "Aktarılacak kayıt yok" in uyarilar[-1]


def test_vazgecilirse_bir_sey_yazilmaz(app, uyarilar, monkeypatch, tablo):
    monkeypatch.setattr(QFileDialog, "getSaveFileName", staticmethod(lambda *a: ("", "")))
    assert da.disa_aktar(None, tablo, "Deneme") is None


def test_kitap_listesi_arama_sonucunu_aktarir(app, uyarilar, monkeypatch, tmp_path):
    lib = Library()
    lib.liste.arama.setText("tahir")
    dosya_sec(monkeypatch, tmp_path / "kitaplar.xlsx")
    da.disa_aktar(lib, lib.liste.tablo, "Kitap Listesi")
    sayfa = load_workbook(tmp_path / "kitaplar.xlsx").active
    assert sayfa.max_row == 3 and sayfa["A1"].value == "Adı" and sayfa.title == "Kitap Listesi"   # gizli Kayıt No yazılmaz


def test_panellerde_aktar_butonlari(app, uyarilar):
    lib = Library()
    q = lib
    assert q.liste.btn_aktar.text() == "Dışa Aktar" and q.liste.tablo.contextMenuPolicy() == Qt.CustomContextMenu
    assert q.istatistik.tablolar[1].contextMenuPolicy() == Qt.CustomContextMenu
    assert lib.gecmis.aktar.text() == "Dışa Aktar" and lib.odunc.btn_aktar.text() == "Dışa Aktar"
    assert lib.filtre.btn_aktar.text() == "Dışa Aktar" and lib.filtre.tablo.contextMenuPolicy() == Qt.CustomContextMenu
    g = Guest()
    assert g.liste.btn_aktar.text() == "Dışa Aktar" and g.filtre.tablo.contextMenuPolicy() == Qt.CustomContextMenu
